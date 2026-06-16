import os
import json
from concurrent.futures import ThreadPoolExecutor, as_completed
import fitz  # PyMuPDF
from docx import Document
import openpyxl
import pytesseract
from PIL import Image
import io

# Tesseract-ის გზა
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
POPPLER_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "poppler", "bin")
os.environ["TESSDATA_PREFIX"] = r"C:\Program Files\Tesseract-OCR\tessdata"

CACHE_FILE = "cache/cache.json"
CACHE_VERSION = 6
OCR_LANG = os.getenv("DOCUMIND_OCR_LANG", "eng+kat")
OCR_CONFIG = "--oem 3 --psm 6"
SKIP_DIRS = {'venv', '__pycache__', '.git'}
OCR_DPI = 220
MIN_TEXT_LEN_SKIP_OCR = 120
MIN_ALNUM_RATIO_SKIP_OCR = 0.20
MAX_WORKERS = max(2, min(8, (os.cpu_count() or 4)))


def should_run_ocr(extracted_text):
    if not extracted_text:
        return True
    normalized = extracted_text.strip()
    if len(normalized) < MIN_TEXT_LEN_SKIP_OCR:
        return True

    alnum_count = sum(ch.isalnum() for ch in normalized)
    ratio = alnum_count / max(len(normalized), 1)
    return ratio < MIN_ALNUM_RATIO_SKIP_OCR


def parse_file(filepath, filename, ext, modified_time):
    if ext == '.pdf':
        text = extract_text_from_pdf(filepath)
    elif ext == '.docx':
        text = extract_text_from_docx(filepath)
    elif ext == '.xlsx':
        text = extract_text_from_xlsx(filepath)
    else:
        text = ""

    doc_data = {
        'filename': filename,
        'filepath': filepath,
        'extension': ext,
        'text': text,
        'text_lower': text.lower(),
        'modified': modified_time
    }
    return filepath, modified_time, doc_data

def extract_text_from_pdf(filepath):
    parts = []
    try:
        doc = fitz.open(filepath)
        for page in doc:
            text = page.get_text().strip()
            if text:
                parts.append(text)

            if should_run_ocr(text):
                print("  OCR scanning page...")
                pix = page.get_pixmap(dpi=OCR_DPI)
                img_data = pix.tobytes("png")
                img = Image.open(io.BytesIO(img_data))

                try:
                    ocr_text = pytesseract.image_to_string(
                        img,
                        lang=OCR_LANG,
                        config=OCR_CONFIG
                    )
                except Exception as e:
                    print(f"  OCR warning (eng+kat failed): {e}")
                    ocr_text = pytesseract.image_to_string(
                        img,
                        lang='eng',
                        config=OCR_CONFIG
                    )

                if ocr_text.strip():
                    parts.append(ocr_text)
        doc.close()
    except Exception as e:
        print(f"PDF error: {filepath} - {e}")
    return "\n".join(parts)

def extract_text_from_docx(filepath):
    parts = []
    try:
        doc = Document(filepath)
        for para in doc.paragraphs:
            parts.append(para.text)
    except Exception as e:
        print(f"DOCX error: {filepath} - {e}")
    return "\n".join(parts)

def extract_text_from_xlsx(filepath):
    parts = []
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for sheet in wb.worksheets:
            headers = []
            for row in sheet.iter_rows(values_only=True):
                if not headers:
                    headers = [str(h).lower() if h else "" for h in row]
                    continue
                row_parts = []
                for i, cell in enumerate(row):
                    if i < len(headers) and cell:
                        if headers[i] == 'deadlines':
                            row_parts.append(f"expires: {cell}")
                        else:
                            row_parts.append(f"{headers[i]}: {cell}")
                if row_parts:
                    parts.append("\n".join(row_parts))
    except Exception as e:
        print(f"XLSX error: {filepath} - {e}")
    return "\n\n".join(parts)

def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            return {}
    return {}

def save_cache(cache):
    try:
        os.makedirs(os.path.dirname(CACHE_FILE), exist_ok=True)
        with open(CACHE_FILE, 'w', encoding='utf-8') as f:
            json.dump(cache, f, ensure_ascii=False)
    except Exception as e:
        print(f"Cache save error: {e}")

def scan_folder(folder_path):
    documents = []
    supported = ('.pdf', '.docx', '.xlsx')
    cache = load_cache()
    updated = False

    if isinstance(folder_path, str):
        roots = [p.strip() for p in folder_path.split(';') if p.strip()]
    else:
        roots = list(folder_path)

    pending = []
    for scan_root in roots:
        if not os.path.isdir(scan_root):
            print(f"Skip path (not found): {scan_root}")
            continue
        for root, dirs, files in os.walk(scan_root):
            dirs[:] = [d for d in dirs if d.lower() not in SKIP_DIRS]

            for filename in files:
                ext = os.path.splitext(filename)[1].lower()
                if ext in supported:
                    filepath = os.path.join(root, filename)
                    modified_time = str(os.path.getmtime(filepath))

                    # Reuse cached document text when file is unchanged.
                    cached = cache.get(filepath, {})
                    if (
                        cached
                        and cached.get('modified') == modified_time
                        and cached.get('cache_version') == CACHE_VERSION
                    ):
                        cached_data = cache[filepath].get('data', {})
                        cached_data.setdefault('filename', filename)
                        cached_data.setdefault('filepath', filepath)
                        cached_data.setdefault('extension', ext)
                        cached_data['modified'] = modified_time
                        documents.append(cached_data)
                        print(f"✓ Cached: {filename}")
                        continue

                    pending.append((filepath, filename, ext, modified_time))

    if pending:
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = [
                executor.submit(parse_file, filepath, filename, ext, modified_time)
                for filepath, filename, ext, modified_time in pending
            ]
            for future in as_completed(futures):
                try:
                    filepath, modified_time, doc_data = future.result()
                except Exception as e:
                    print(f"Parse error: {e}")
                    continue

                cache[filepath] = {
                    'modified': modified_time,
                    'cache_version': CACHE_VERSION,
                    'data': doc_data
                }
                updated = True
                documents.append(doc_data)
                print(f"✓ Scanned: {doc_data['filename']}")

    if updated:
        save_cache(cache)

    return documents